from datetime import datetime, timezone
from io import BytesIO

from fastapi import (
    APIRouter,
    Depends,
    HTTPException,
    status,
)

from fastapi.responses import StreamingResponse

from pydantic import BaseModel

from sqlalchemy.orm import Session

from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)

from docx import Document

from app.core.dependencies import get_current_user
from app.database.connection import get_db

from app.models import (
    Application,
    Internship,
    User,
)

# =========================================================
# NOTIFICATION SERVICE
# =========================================================

from app.services.notification_service import (
    create_user_notification,
)


router = APIRouter(
    prefix="/api/applications",
    tags=["Applications"],
)


# =========================================================
# APPLY REQUEST
# =========================================================

class ApplyInternshipRequest(BaseModel):
    cover_letter: str


# =========================================================
# GET MY APPLICATIONS
# =========================================================

@router.get("/my")
def get_my_applications(
    db: Session = Depends(get_db),
    current_user: User = Depends(
        get_current_user
    ),
):

    # =====================================================
    # CHECK ROLE
    # =====================================================

    if current_user.role != "intern":

        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=(
                "Only interns can view "
                "their applications"
            ),
        )

    # =====================================================
    # GET APPLICATIONS
    # =====================================================

    applications = (
        db.query(Application)
        .join(
            Internship,
            Internship.id
            == Application.internship_id,
        )
        .filter(
            Application.user_id
            == current_user.id,
        )
        .order_by(
            Application.applied_at.desc()
        )
        .all()
    )

    result = []

    # =====================================================
    # BUILD RESPONSE
    # =====================================================

    for application in applications:

        internship = application.internship

        can_withdraw = (
            application.status == "pending"
        )

        result.append(
            {
                "application_id": (
                    application.id
                ),

                "internship_id": (
                    internship.id
                ),

                "company_name": (
                    internship.company_name
                ),

                "job_title": (
                    internship.title
                ),

                "location": (
                    internship.location
                ),

                "duration": (
                    internship.duration
                ),

                "work_mode": (
                    internship.work_mode
                ),

                "stipend": (
                    internship.stipend
                ),

                "cover_letter": (
                    application.cover_letter
                ),

                "status": (
                    application.status
                ),

                "applied_at": (
                    application.applied_at
                ),

                "withdrawn_at": (
                    application.withdrawn_at
                ),

                "can_withdraw": (
                    can_withdraw
                ),
            }
        )

    # =====================================================
    # RESPONSE
    # =====================================================

    return {
        "message": (
            "Applications retrieved successfully"
        ),

        "total_applications": len(
            result
        ),

        "applications": result,
    }


# =========================================================
# GET APPLICATION DETAILS
# =========================================================

@router.get("/{application_id}")
def get_application_details(
    application_id: int,

    db: Session = Depends(get_db),

    current_user: User = Depends(
        get_current_user
    ),
):

    # =====================================================
    # CHECK ROLE
    # =====================================================

    if current_user.role != "intern":

        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=(
                "Only interns can view "
                "application details"
            ),
        )

    # =====================================================
    # FIND APPLICATION
    # =====================================================

    application = (
        db.query(Application)
        .join(
            Internship,
            Internship.id
            == Application.internship_id,
        )
        .filter(
            Application.id
            == application_id,

            Application.user_id
            == current_user.id,
        )
        .first()
    )

    if not application:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Application not found",
        )

    # =====================================================
    # INTERNSHIP
    # =====================================================

    internship = application.internship

    can_withdraw = (
        application.status == "pending"
    )

    # =====================================================
    # RESPONSE
    # =====================================================

    return {
        "message": (
            "Application details retrieved successfully"
        ),

        "application": {

            "application_id": (
                application.id
            ),

            "status": (
                application.status
            ),

            "applied_at": (
                application.applied_at
            ),

            "withdrawn_at": (
                application.withdrawn_at
            ),

            "can_withdraw": (
                can_withdraw
            ),

            "cover_letter": (
                application.cover_letter
            ),

            "internship": {

                "id": internship.id,

                "company_name": (
                    internship.company_name
                ),

                "title": (
                    internship.title
                ),

                "description": (
                    internship.description
                ),

                "location": (
                    internship.location
                ),

                "duration": (
                    internship.duration
                ),

                "work_mode": (
                    internship.work_mode
                ),

                "stipend": (
                    internship.stipend
                ),

                "start_date": (
                    internship.start_date
                ),

                "required_skills": (
                    internship.required_skills
                ),

                "eligibility": (
                    internship.eligibility
                ),

                "responsibilities": (
                    internship.responsibilities
                ),

                "benefits": (
                    internship.benefits
                ),

                "application_url": (
                    internship.application_url
                ),
            },
        },
    }


# =========================================================
# APPLY FOR INTERNSHIP
# =========================================================

@router.post(
    "/{internship_id}",
    status_code=status.HTTP_201_CREATED,
)
def apply_for_internship(
    internship_id: int,

    request: ApplyInternshipRequest,

    db: Session = Depends(get_db),

    current_user: User = Depends(
        get_current_user
    ),
):

    # =====================================================
    # CHECK ROLE
    # =====================================================

    if current_user.role != "intern":

        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=(
                "Only interns can apply "
                "for internships"
            ),
        )

    # =====================================================
    # VALIDATE COVER LETTER
    # =====================================================

    cover_letter = (
        request.cover_letter or ""
    ).strip()

    if not cover_letter:

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Cover letter is required",
        )

    # =====================================================
    # FIND INTERNSHIP
    # =====================================================

    internship = (
        db.query(Internship)
        .filter(
            Internship.id == internship_id
        )
        .first()
    )

    if not internship:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Internship not found",
        )

    # =====================================================
    # CHECK EXISTING APPLICATION
    # =====================================================

    existing_application = (
        db.query(Application)
        .filter(
            Application.user_id
            == current_user.id,

            Application.internship_id
            == internship_id,
        )
        .first()
    )

    if existing_application:

        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "You have already applied "
                "for this internship"
            ),
        )

    # =====================================================
    # CREATE APPLICATION
    # =====================================================

    application = Application(
        user_id=current_user.id,

        internship_id=internship.id,

        cover_letter=cover_letter,

        status="pending",
    )

    db.add(application)

    db.commit()

    db.refresh(application)

    # =====================================================
    # INTERN NOTIFICATION
    # =====================================================

    create_user_notification(
        db=db,

        user_id=current_user.id,

        title="Application Submitted",

        message=(
            f"Your application for "
            f"{internship.title} at "
            f"{internship.company_name} "
            "has been submitted successfully."
        ),

        notification_type=(
            "application_submitted"
        ),

        related_entity_type=(
            "application"
        ),

        related_entity_id=(
            application.id
        ),
    )

    # =====================================================
    # RESPONSE
    # =====================================================

    return {
        "message": (
            "Internship application submitted "
            "successfully"
        ),

        "application": {

            "id": (
                application.id
            ),

            "internship_id": (
                application.internship_id
            ),

            "company_name": (
                internship.company_name
            ),

            "job_title": (
                internship.title
            ),

            "cover_letter": (
                application.cover_letter
            ),

            "status": (
                application.status
            ),

            "applied_at": (
                application.applied_at
            ),

            "withdrawn_at": (
                application.withdrawn_at
            ),
        },
    }


# =========================================================
# WITHDRAW APPLICATION
# =========================================================

@router.patch(
    "/{application_id}/withdraw"
)
def withdraw_application(
    application_id: int,

    db: Session = Depends(get_db),

    current_user: User = Depends(
        get_current_user
    ),
):

    # =====================================================
    # CHECK ROLE
    # =====================================================

    if current_user.role != "intern":

        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=(
                "Only interns can withdraw "
                "applications"
            ),
        )

    # =====================================================
    # FIND APPLICATION
    # =====================================================

    application = (
        db.query(Application)
        .filter(
            Application.id
            == application_id,

            Application.user_id
            == current_user.id,
        )
        .first()
    )

    if not application:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Application not found",
        )

    # =====================================================
    # CHECK STATUS
    # =====================================================

    if application.status != "pending":

        if application.status == "approved":

            detail = (
                "Approved applications "
                "cannot be withdrawn"
            )

        elif application.status == "rejected":

            detail = (
                "Rejected applications "
                "cannot be withdrawn"
            )

        elif application.status == "withdrawn":

            detail = (
                "Application has already "
                "been withdrawn"
            )

        else:

            detail = (
                "This application cannot "
                "be withdrawn"
            )

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=detail,
        )

    # =====================================================
    # CURRENT TIME
    # =====================================================

    current_time = (
        datetime.now(
            timezone.utc
        ).replace(
            tzinfo=None
        )
    )

    # =====================================================
    # UPDATE APPLICATION
    # =====================================================

    application.status = "withdrawn"

    application.withdrawn_at = (
        current_time
    )

    application.updated_at = (
        current_time
    )

    db.commit()

    db.refresh(application)

    # =====================================================
    # FIND INTERNSHIP
    # =====================================================

    internship = (
        db.query(Internship)
        .filter(
            Internship.id
            == application.internship_id
        )
        .first()
    )

    # =====================================================
    # INTERN NOTIFICATION
    # =====================================================

    create_user_notification(
        db=db,

        user_id=current_user.id,

        title="Application Withdrawn",

        message=(
            f"Your application for "
            f"{internship.title if internship else 'the internship'} "
            "has been withdrawn successfully."
        ),

        notification_type=(
            "application_withdrawn"
        ),

        related_entity_type=(
            "application"
        ),

        related_entity_id=(
            application.id
        ),
    )

    # =====================================================
    # RESPONSE
    # =====================================================

    return {
        "message": (
            "Application withdrawn successfully"
        ),

        "application": {

            "id": (
                application.id
            ),

            "internship_id": (
                application.internship_id
            ),

            "company_name": (
                internship.company_name
                if internship
                else None
            ),

            "job_title": (
                internship.title
                if internship
                else None
            ),

            "status": (
                application.status
            ),

            "applied_at": (
                application.applied_at
            ),

            "withdrawn_at": (
                application.withdrawn_at
            ),

            "can_withdraw": False,
        },
    }


# =========================================================
# DOWNLOAD COVER LETTER PDF
# =========================================================

@router.get(
    "/{application_id}/cover-letter/pdf"
)
def download_cover_letter_pdf(
    application_id: int,

    db: Session = Depends(get_db),

    current_user: User = Depends(
        get_current_user
    ),
):

    # =====================================================
    # CHECK ROLE
    # =====================================================

    if current_user.role != "intern":

        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=(
                "Only interns can download "
                "cover letters"
            ),
        )

    # =====================================================
    # FIND APPLICATION
    # =====================================================

    application = (
        db.query(Application)
        .filter(
            Application.id
            == application_id,

            Application.user_id
            == current_user.id,
        )
        .first()
    )

    if not application:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Application not found",
        )

    # =====================================================
    # CHECK COVER LETTER
    # =====================================================

    if not application.cover_letter:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Cover letter not found",
        )

    # =====================================================
    # FIND INTERNSHIP
    # =====================================================

    internship = (
        db.query(Internship)
        .filter(
            Internship.id
            == application.internship_id
        )
        .first()
    )

    # =====================================================
    # CREATE PDF
    # =====================================================

    buffer = BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=50,
        leftMargin=50,
        topMargin=50,
        bottomMargin=50,
    )

    styles = getSampleStyleSheet()

    normal_style = styles["Normal"]

    normal_style.alignment = TA_LEFT

    normal_style.fontSize = 11

    normal_style.leading = 17

    story = [

        Paragraph(
            "Cover Letter",
            styles["Title"],
        ),

        Spacer(
            1,
            20,
        ),
    ]

    # =====================================================
    # ADD COVER LETTER PARAGRAPHS
    # =====================================================

    for paragraph in application.cover_letter.split(
        "\n"
    ):

        paragraph = paragraph.strip()

        if paragraph:

            paragraph = (
                paragraph
                .replace(
                    "&",
                    "&amp;",
                )
                .replace(
                    "<",
                    "&lt;",
                )
                .replace(
                    ">",
                    "&gt;",
                )
            )

            story.append(
                Paragraph(
                    paragraph,
                    normal_style,
                )
            )

            story.append(
                Spacer(
                    1,
                    8,
                )
            )

        else:

            story.append(
                Spacer(
                    1,
                    6,
                )
            )

    # =====================================================
    # BUILD DOCUMENT
    # =====================================================

    document.build(
        story
    )

    buffer.seek(0)

    company_name = (
        internship.company_name
        if internship
        else "Company"
    )

    filename = (
        f"{company_name}_Cover_Letter.pdf"
    )

    # =====================================================
    # RESPONSE
    # =====================================================

    return StreamingResponse(
        buffer,

        media_type="application/pdf",

        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            )
        },
    )


# =========================================================
# DOWNLOAD COVER LETTER DOCX
# =========================================================

@router.get(
    "/{application_id}/cover-letter/docx"
)
def download_cover_letter_docx(
    application_id: int,

    db: Session = Depends(get_db),

    current_user: User = Depends(
        get_current_user
    ),
):

    # =====================================================
    # CHECK ROLE
    # =====================================================

    if current_user.role != "intern":

        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=(
                "Only interns can download "
                "cover letters"
            ),
        )

    # =====================================================
    # FIND APPLICATION
    # =====================================================

    application = (
        db.query(Application)
        .filter(
            Application.id
            == application_id,

            Application.user_id
            == current_user.id,
        )
        .first()
    )

    if not application:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Application not found",
        )

    # =====================================================
    # CHECK COVER LETTER
    # =====================================================

    if not application.cover_letter:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Cover letter not found",
        )

    # =====================================================
    # FIND INTERNSHIP
    # =====================================================

    internship = (
        db.query(Internship)
        .filter(
            Internship.id
            == application.internship_id
        )
        .first()
    )

    # =====================================================
    # CREATE DOCX
    # =====================================================

    document = Document()

    document.add_heading(
        "Cover Letter",
        level=1,
    )

    # =====================================================
    # ADD COVER LETTER PARAGRAPHS
    # =====================================================

    for paragraph in application.cover_letter.split(
        "\n"
    ):

        paragraph = paragraph.strip()

        if paragraph:

            document.add_paragraph(
                paragraph
            )

        else:

            document.add_paragraph(
                ""
            )

    # =====================================================
    # CREATE BUFFER
    # =====================================================

    buffer = BytesIO()

    document.save(
        buffer
    )

    buffer.seek(0)

    company_name = (
        internship.company_name
        if internship
        else "Company"
    )

    filename = (
        f"{company_name}_Cover_Letter.docx"
    )

    # =====================================================
    # RESPONSE
    # =====================================================

    return StreamingResponse(
        buffer,

        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.wordprocessingml.document"
        ),

        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            )
        },
    )